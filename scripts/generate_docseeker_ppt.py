#!/usr/bin/env python3
"""Generate a compact Chinese academic PPT for arXiv:2604.12812.

The script intentionally downloads the paper PDF at runtime and crops only
figure/table regions (never full pages) for insertion into the deck.
"""
from __future__ import annotations

import argparse
import io
import urllib.request
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

PAPER_URL = "https://arxiv.org/pdf/2604.12812"
OUT_DIR = Path("outputs/docseeker")
PDF_PATH = OUT_DIR / "docseeker_2604.12812.pdf"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "DocSeeker_Chinese_Academic_PPT.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TITLE_C = RGBColor(24, 55, 92)
ACCENT = RGBColor(36, 128, 155)
GREEN = RGBColor(53, 145, 78)
RED = RGBColor(187, 67, 58)
DARK = RGBColor(35, 42, 50)
LIGHT_BG = RGBColor(246, 249, 252)
MUTED = RGBColor(94, 108, 122)


def download_pdf(force: bool = False) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if PDF_PATH.exists() and not force:
        return PDF_PATH
    req = urllib.request.Request(PAPER_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=90) as resp:
        PDF_PATH.write_bytes(resp.read())
    return PDF_PATH


def crop_pdf_region(doc: fitz.Document, page_index: int, box: tuple[float, float, float, float], name: str, zoom: float = 3.0) -> Path:
    """Crop by normalized coordinates (x0, y0, x1, y1) and save as PNG."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    page = doc[page_index]
    rect = page.rect
    clip = fitz.Rect(
        rect.x0 + box[0] * rect.width,
        rect.y0 + box[1] * rect.height,
        rect.x0 + box[2] * rect.width,
        rect.y0 + box[3] * rect.height,
    )
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=clip, alpha=False)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    path = ASSET_DIR / f"{name}.png"
    img.save(path, optimize=True)
    return path


def set_text_frame(tf, font_size=18, color=DARK, bold=False, font="Microsoft YaHei"):
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.bold = bold


def add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.20), Inches(12.45), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TITLE_C
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.50), Inches(0.72), Inches(12.3), Inches(0.28))
        stf = sub.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.name = "Microsoft YaHei"
        sr.font.size = Pt(10.5)
        sr.font.color.rgb = MUTED
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.96), Inches(12.45), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_bullets(slide, bullets: list[str], x, y, w, h, font_size=15.5, color=DARK, bullet=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, txt in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
        if bullet:
            p.text = "• " + txt
    return box


def add_section_label(slide, text, x, y, w, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    return shape


def add_card(slide, x, y, w, h, title, body, fill=LIGHT_BG, accent=ACCENT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = RGBColor(214, 224, 235)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    tb = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), w - Inches(0.3), h - Inches(0.2))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = TITLE_C
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.name = "Microsoft YaHei"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = DARK
    p2.space_before = Pt(5)
    return card


def add_picture_fit(slide, img_path: Path, x, y, w, h):
    with Image.open(img_path) as im:
        iw, ih = im.size
    ratio = min(w / iw, h / ih)
    pw, ph = int(iw * ratio), int(ih * ratio)
    left = x + (w - pw) / 2
    top = y + (h - ph) / 2
    return slide.shapes.add_picture(str(img_path), left, top, width=pw, height=ph)


def add_caption(slide, text, x, y, w):
    box = slide.shapes.add_textbox(x, y, w, Inches(0.26))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED


def make_deck(pdf: Path = PDF_PATH, out: Path = PPTX_PATH) -> Path:
    doc = fitz.open(pdf)
    assets = {
        "fig1": crop_pdf_region(doc, 1, (0.10, 0.07, 0.89, 0.40), "fig1_overview_alr"),
        "fig2": crop_pdf_region(doc, 3, (0.08, 0.04, 0.92, 0.41), "fig2_training_framework"),
        "tab1": crop_pdf_region(doc, 5, (0.06, 0.05, 0.95, 0.31), "table1_main_results"),
        "fig3": crop_pdf_region(doc, 6, (0.10, 0.05, 0.54, 0.30), "fig3_length_robustness"),
        "fig4": crop_pdf_region(doc, 6, (0.12, 0.37, 0.88, 0.55), "fig4_rag_synergy"),
        "tab3": crop_pdf_region(doc, 6, (0.08, 0.55, 0.92, 0.78), "table3_data_ablation"),
        "tab45": crop_pdf_region(doc, 7, (0.07, 0.06, 0.93, 0.28), "table4_5_ablation"),
        "tab7": crop_pdf_region(doc, 12, (0.06, 0.83, 0.45, 0.93), "table7_latency"),
    }

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # Slide 1
    s = prs.slides.add_slide(blank)
    add_title(s, "DocSeeker：面向长文档理解的结构化视觉推理与证据定位", "CVPR 2026 Highlight · Hao Yan et al. · arXiv:2604.12812v5")
    add_bullets(s, [
        "核心问题：长文档中关键信息被大量无关页面淹没，且训练集通常只有短答案监督。",
        "核心范式：让 MLLM 显式执行 Analysis → Localization → Reasoning，并输出 evidence_pages。",
        "训练框架：ALR CoT 蒸馏 SFT + Evidence-aware GRPO + EGRA 分辨率分配。",
        "主要结果：7B 开源模型在 5 个文档 VQA 基准上显著超过同架构 baseline，并具备强 OOD/超长文档泛化。",
    ], Inches(0.65), Inches(1.25), Inches(5.4), Inches(2.55), 16.5)
    add_picture_fit(s, assets["fig1"], Inches(6.20), Inches(1.15), Inches(6.55), Inches(3.25))
    add_caption(s, "原论文 Figure 1（紧裁）：主要实验现象 + ALR 工作流", Inches(6.20), Inches(4.35), Inches(6.55))
    add_bullets(s, [
        "随页数增加，DocSeeker 的性能下降明显小于 baseline。",
        "ALR 输出把“答案”与“证据页”绑定，可验证性更强。",
        "跨 DUDE/MP-DocVQA/MMLongBench-doc/LongDocURL/SlideVQA 均有稳定增益。",
    ], Inches(6.35), Inches(4.72), Inches(6.15), Inches(1.45), 13.2)

    # Slide 2
    s = prs.slides.add_slide(blank)
    add_title(s, "背景与动机：长文档 VQA 的两个瓶颈", "纯视觉 MLLM 保留版面信息，但在长上下文中面临低 SNR 与弱监督。")
    add_card(s, Inches(0.65), Inches(1.25), Inches(3.85), Inches(1.65), "瓶颈 1：低信噪比", "证据页通常只占少数；Top-K RAG 面临“召回 vs 噪声”两难。", accent=RED)
    add_card(s, Inches(4.75), Inches(1.25), Inches(3.85), Inches(1.65), "瓶颈 2：监督稀缺", "多数数据集只给最终短答案，缺少证据定位与推理路径监督。", accent=RED)
    add_card(s, Inches(8.85), Inches(1.25), Inches(3.85), Inches(1.65), "目标：可验证泛化", "模型不仅答对，还要说明“在哪些页上、如何推理得到答案”。", accent=GREEN)
    add_section_label(s, "本文核心主张", Inches(0.70), Inches(3.35), Inches(1.60))
    add_bullets(s, [
        "把长文档问答从直接生成短答案，改造成带页面锚点的结构化推理任务。",
        "利用 teacher 仅基于最小证据上下文蒸馏高质量 ALR CoT，避免昂贵/噪声大的全页蒸馏。",
        "在 RL 阶段同时优化格式、证据定位和答案正确性，让定位能力服务于最终回答。",
        "训练时用 EGRA 对证据页保真、对非证据页降分辨率，在固定 token 预算下保留上下文。",
    ], Inches(0.90), Inches(3.90), Inches(11.8), Inches(2.25), 17)

    # Slide 3
    s = prs.slides.add_slide(blank)
    add_title(s, "方法：ALR 范式 + 两阶段训练 + EGRA", "方法的关键不是更大模型，而是把“定位证据”变成可学习、可奖励的中间能力。")
    add_picture_fit(s, assets["fig2"], Inches(0.55), Inches(1.18), Inches(7.05), Inches(3.95))
    add_caption(s, "原论文 Figure 2（紧裁）：训练框架与 Evidence-Guided Resolution Allocation", Inches(0.55), Inches(5.02), Inches(7.05))
    add_bullets(s, [
        "Stage I：用 Gemini-2.5-Flash 蒸馏 ALR CoT；EM + GPT-4o 二次校验保证数据质量。",
        "Stage II：EviGRPO 奖励 = 格式 + 证据页定位 + 答案正确性，避免只优化短答案。",
        "Page ID 作为视觉 token 的页面指针，使定位输出和输入页面显式对齐。",
        "EGRA：证据页高分辨率；非证据页低分辨率并保留少量高分辨率干扰页，提高鲁棒性。",
    ], Inches(7.85), Inches(1.25), Inches(4.75), Inches(3.6), 14.2)

    # Slide 4
    s = prs.slides.add_slide(blank)
    add_title(s, "主实验：5 个基准上的整体性能", "DocSeeker 不依赖检索，作为 7B 纯视觉模型在 in-domain 与 OOD 基准上均有强竞争力。")
    add_picture_fit(s, assets["tab1"], Inches(0.55), Inches(1.10), Inches(12.25), Inches(2.85))
    add_caption(s, "原论文 Table 1（紧裁）：DUDE、MPDocVQA、MMLongBench-doc、LongDocURL、SlideVQA", Inches(0.55), Inches(3.87), Inches(12.25))
    add_bullets(s, [
        "相对同架构 Baseline：DUDE 35.2→57.4，MPDocVQA 70.1→86.2，MMLong 25.4→40.1。",
        "短答案 SFT 主要提升 in-domain；ALR SFT 显著提升 OOD，说明结构化证据监督带来泛化。",
        "EviGRPO 在 SFT 基础上进一步提升全部 5 个基准，定位奖励与答案奖励互补。",
        "与 GPT-4o 等闭源模型相比，DocSeeker 在多个 OOD 指标上接近或超过，且输出证据页。",
    ], Inches(0.85), Inches(4.35), Inches(11.7), Inches(1.85), 14.5)

    # Slide 5
    s = prs.slides.add_slide(blank)
    add_title(s, "分析：长文档鲁棒性与 RAG 协同", "ALR 的价值体现在能在噪声页中稳定定位，而不是只依赖检索器提前找准证据。")
    add_picture_fit(s, assets["fig3"], Inches(0.65), Inches(1.18), Inches(5.55), Inches(2.15))
    add_picture_fit(s, assets["fig4"], Inches(6.75), Inches(1.18), Inches(5.55), Inches(2.15))
    add_caption(s, "原论文 Figure 3/4（紧裁）：文档长度影响与 ColQwen2.5 检索集成", Inches(0.65), Inches(3.35), Inches(11.65))
    add_bullets(s, [
        "长度分析：Baseline 随上下文从 10→60 页降至 13.9；DocSeeker 基本稳定在约 30+。",
        "全页 vs 证据页：Baseline 全文输入大幅掉点；DocSeeker 全文输入几乎接近 evidence-only 上界。",
        "RAG 分析：K 增大时检索召回提高但噪声上升；DocSeeker 能在 noisy Top-K 中继续细粒度定位。",
        "启示：检索器负责粗筛，DocSeeker 负责带证据页的精读推理，可缓解 Top-K dilemma。",
    ], Inches(0.90), Inches(4.05), Inches(11.6), Inches(2.0), 14.5)

    # Slide 6
    s = prs.slides.add_slide(blank)
    add_title(s, "消融与效率：哪些设计真正起作用？", "消融验证 ALR 数据、Page ID、EGRA 与 evidence-aware RL 都是关键组件。")
    add_picture_fit(s, assets["tab3"], Inches(0.55), Inches(1.15), Inches(6.0), Inches(2.35))
    add_picture_fit(s, assets["tab45"], Inches(6.75), Inches(1.15), Inches(5.65), Inches(2.35))
    add_picture_fit(s, assets["tab7"], Inches(0.70), Inches(4.72), Inches(3.9), Inches(0.95))
    add_caption(s, "原论文 Table 3/4/5/7（紧裁）：数据、分辨率、RL 奖励与延迟消融", Inches(0.55), Inches(3.55), Inches(11.85))
    add_bullets(s, [
        "数据范式：ALR CoT > Vanilla CoT > raw short-answer；去掉 Page ID 明显降低 Acc/F1。",
        "数据规模：ALR CoT 从 20% 到全量，MMLong Acc 32.7→38.6，说明结构化数据可扩展。",
        "EGRA 优于固定分辨率/截断；保留低分辨率非证据页比直接丢弃更好。",
        "EviGRPO 最佳权重提升至 40.1/38.4；延迟从 19s 到 25s，换来 +14.7 Acc 与可解释性。",
    ], Inches(4.85), Inches(4.20), Inches(7.55), Inches(1.85), 13.6)

    # Slide 7
    s = prs.slides.add_slide(blank)
    add_title(s, "结论与讨论", "DocSeeker 将长文档理解从“读完整本后猜答案”转为“先定位证据、再基于证据推理”。")
    add_section_label(s, "主要结论", Inches(0.75), Inches(1.30), Inches(1.45), GREEN)
    add_bullets(s, [
        "ALR 让模型学习可迁移的页面级证据定位能力，是 OOD/超长文档泛化的核心。",
        "EviGRPO 将定位质量显式纳入奖励，避免 RL 只追求最终答案而忽略可验证性。",
        "EGRA 在长视觉序列训练中兼顾上下文覆盖与关键页清晰度，降低 token/显存压力。",
        "与 RAG 自然互补：检索降低候选空间，DocSeeker 在候选页内抵抗噪声并给出证据页。",
    ], Inches(0.95), Inches(1.86), Inches(5.65), Inches(2.35), 15.3)
    add_section_label(s, "局限与后续方向", Inches(7.05), Inches(1.30), Inches(1.85), RED)
    add_bullets(s, [
        "结构化输出增加推理 token；论文显示端到端延迟只中等增加，但高吞吐场景仍需优化。",
        "证据页监督依赖已有数据与 teacher/judge 质量，复杂跨页证据的标注噪声仍可能影响训练。",
        "未来可探索更强检索器、更细粒度区域证据、以及端到端可验证的多文档推理系统。",
    ], Inches(7.20), Inches(1.86), Inches(5.35), Inches(2.0), 15.3)
    add_card(s, Inches(1.05), Inches(5.10), Inches(11.2), Inches(1.05), "一句话 takeaway", "DocSeeker 的贡献不是单点技巧，而是把“证据定位”贯穿输入表示、SFT 数据、RL 奖励和分辨率分配，从而系统性解决长文档低 SNR 与弱监督。", fill=RGBColor(232, 244, 248), accent=ACCENT)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def validate_pptx(path: Path, expected_slides: int = 7) -> None:
    prs = Presentation(path)
    assert len(prs.slides) == expected_slides, f"expected {expected_slides} slides, got {len(prs.slides)}"
    assert path.stat().st_size > 100_000, f"pptx too small: {path.stat().st_size} bytes"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force-download", action="store_true")
    args = parser.parse_args()
    pdf = download_pdf(force=args.force_download)
    out = make_deck(pdf)
    validate_pptx(out)
    print(f"Generated: {out}")
    print("Slide count: 7")


if __name__ == "__main__":
    main()
