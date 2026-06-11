#!/usr/bin/env python3
"""Validate the generated DocSeeker deck.

This script performs offline structural validation. If python-pptx is installed,
it also reopens the file with Presentation(); otherwise it reports that the
optional dependency is unavailable while keeping the deterministic zip/XML checks.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

PPTX = Path("outputs/docseeker_compact_ppt.pptx")
BANNED = ["紧裁", "裁剪", "截图", "cropped", "tight crop", "extracted from page"]


def main() -> int:
    if not PPTX.exists():
        print(f"FAIL: missing {PPTX}")
        return 1
    with zipfile.ZipFile(PPTX) as z:
        bad = z.testzip()
        if bad:
            print(f"FAIL: corrupted zip member {bad}")
            return 1
        slides = sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        if len(slides) > 7:
            print(f"FAIL: slide count {len(slides)} exceeds 7")
            return 1
        text = "\n".join(z.read(s).decode("utf-8", errors="ignore") for s in slides)
        for word in BANNED:
            if word.lower() in text.lower():
                print(f"FAIL: banned visible wording found: {word}")
                return 1
        for slide in slides:
            ET.fromstring(z.read(slide))
    print(f"PASS: zip/XML validation succeeded; slide_count={len(slides)}; captions clean")
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        print("WARN: python-pptx is not installed in this environment; CI workflow installs it before artifact validation.")
        return 0
    prs = Presentation(str(PPTX))
    if len(prs.slides) > 7:
        print(f"FAIL: python-pptx slide count {len(prs.slides)} exceeds 7")
        return 1
    print(f"PASS: python-pptx reopened deck; slide_count={len(prs.slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
